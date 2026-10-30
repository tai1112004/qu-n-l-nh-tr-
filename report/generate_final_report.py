# -*- coding: utf-8 -*-
import sys
import os

from report_builder import create_report_document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def build_complete_report():
    output_docx = os.path.abspath("Bao_Cao_Do_An_Tot_Nghiep_Quan_Ly_Phong_Tro.docx")
    doc, add_heading, add_p, add_image_placeholder, add_styled_table = create_report_document(output_docx)

    # =========================================================================
    # 1. TRANG BÌA (COVER PAGE)
    # =========================================================================
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.space_before = Pt(12)
    p_univ.paragraph_format.space_after = Pt(2)
    r = p_univ.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True
    r = p_univ.add_run("TRƯỜNG ĐẠI HỌC ................................\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.bold = True
    r = p_univ.add_run("KHOA CÔNG NGHỆ THÔNG TIN\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True

    # Line divider
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(36)
    r = p_div.add_run("--------------------***--------------------")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    # Logo Placeholder
    table_logo = doc.add_table(rows=1, cols=1)
    table_logo.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = table_logo.cell(0, 0)
    c.width = Inches(2.5)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
    c._tc.get_or_add_tcPr().append(shd)
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
        </w:tcBorders>
    ''')
    c._tc.get_or_add_tcPr().append(borders)
    p_logo = c.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(24)
    p_logo.paragraph_format.space_after = Pt(24)
    r_logo = p_logo.add_run("[ LOGO TRƯỜNG ]\n(Dán ảnh logo vào đây)")
    r_logo.font.name = "Times New Roman"
    r_logo.font.size = Pt(10)
    r_logo.font.color.rgb = RGBColor(100, 116, 139)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(12)
    r = p_title.add_run("ĐỒ ÁN TỐT NGHIỆP ĐẠI HỌC\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)

    r = p_title.add_run("NGÀNH: CÔNG NGHỆ THÔNG TIN\n\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True

    r = p_title.add_run("ĐỀ TÀI:\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.bold = True

    r = p_title.add_run("NGHIÊN CỨU XÂY DỰNG HỆ THỐNG WEBSITE\nQUẢN LÝ VÀ CHO THUÊ NHÀ TRỌ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(18)
    r.bold = True
    r.font.color.rgb = RGBColor(180, 0, 0)

    # Student & Teacher info
    p_info = doc.add_paragraph()
    p_info.paragraph_format.left_indent = Inches(1.5)
    p_info.paragraph_format.space_before = Pt(48)
    p_info.paragraph_format.space_after = Pt(48)
    p_info.paragraph_format.line_spacing = 1.3
    
    r = p_info.add_run("Sinh viên thực hiện:\t")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True
    r = p_info.add_run("...........................................................\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)

    r = p_info.add_run("Mã số sinh viên:\t\t")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True
    r = p_info.add_run("...........................................................\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)

    r = p_info.add_run("Lớp:\t\t\t")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True
    r = p_info.add_run("...........................................................\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)

    r = p_info.add_run("Giảng viên hướng dẫn:\t")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True
    r = p_info.add_run("...........................................................\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)

    # Footer Year
    p_year = doc.add_paragraph()
    p_year.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_year.paragraph_format.space_before = Pt(36)
    r = p_year.add_run("Hà Nội, Năm 2026")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True

    doc.add_page_break()

    # =========================================================================
    # 2. LỜI CẢM ƠN
    # =========================================================================
    add_heading("LỜI CẢM ƠN", level=1, bookmark="sec_thanks")
    add_p("Để hoàn thành đồ án tốt nghiệp này, trước hết em xin gửi lời cảm ơn chân thành và sâu sắc nhất đến quý Thầy/Cô giảng viên hướng dẫn, người đã luôn tận tình chỉ bảo, định hướng và động viên em trong suốt quá trình nghiên cứu và thực hiện đề tài.")
    add_p("Trong thời gian được Thầy/Cô hướng dẫn, em không chỉ tiếp thu thêm nhiều kiến thức chuyên môn quý báu về kiến trúc phần mềm, công nghệ Full-stack hiện đại mà còn rèn luyện được tác phong làm việc nghiêm túc, tính kỷ luật và phương pháp giải quyết vấn đề một cách khoa học. Đây là hành trang vô cùng quan trọng đối với bản thân em trong sự nghiệp sau này.")
    add_p("Em cũng xin trân trọng gửi lời cảm ơn tới toàn thể các Thầy/Cô trong Khoa Công nghệ Thông tin, những người đã tâm huyết giảng dạy, truyền đạt tri thức và tạo mọi điều kiện thuận lợi nhất cho em trong suốt những năm tháng học tập và rèn luyện trên giảng đường đại học.")
    add_p("Cuối cùng, em xin bày tỏ lòng biết ơn sâu sắc đến gia đình, bạn bè đã luôn ở bên động viên, ủng hộ và giúp đỡ em trong suốt quá trình học tập để hoàn thành tốt đề tài này.")
    add_p("Kính chúc quý Thầy/Cô luôn dồi dào sức khỏe, hạnh phúc và gặt hái được nhiều thành công hơn nữa trong sự nghiệp giáo dục cao quý!")
    
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.space_before = Pt(24)
    r = p_sign.add_run("Hà Nội, ngày ... tháng ... năm 2026\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.italic = True
    r = p_sign.add_run("Sinh viên thực hiện đề tài\n\n\n\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.bold = True
    r = p_sign.add_run("(Ký và ghi rõ họ tên)")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.italic = True

    doc.add_page_break()

    # =========================================================================
    # 3. MỤC LỤC (INTERACTIVE TABLE OF CONTENTS)
    # =========================================================================
    p_toc_head = doc.add_paragraph()
    p_toc_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_head.paragraph_format.space_before = Pt(12)
    p_toc_head.paragraph_format.space_after = Pt(12)
    r = p_toc_head.add_run("MỤC LỤC")
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)

    p_note = doc.add_paragraph()
    p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_note.paragraph_format.space_after = Pt(12)
    r_note = p_note.add_run("(Nhấn giữ Ctrl và Click vào mục bất kỳ để di chuyển tới nội dung. Khi cần cập nhật số trang, click chuột phải vào mục lục chọn 'Update Field')")
    r_note.font.name = "Times New Roman"
    r_note.font.size = Pt(10.5)
    r_note.italic = True
    r_note.font.color.rgb = RGBColor(100, 116, 139)

    # Word TOC Field Definition
    p_toc = doc.add_paragraph()
    r1 = p_toc.add_run()
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    r1._r.append(fld1)
    
    r2 = p_toc.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    r2._r.append(instr)
    
    r3 = p_toc.add_run()
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    r3._r.append(fld2)

    # Helper function to generate clickable pre-populated TOC item
    def add_toc_row(title, bookmark, level=1, page=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25 * (level - 1))
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        pPr = p._p.get_or_add_pPr()
        tabs = parse_xml(f'''
            <w:tabs {nsdecls("w")}>
                <w:tab w:val="right" w:leader="dot" w:pos="9350"/>
            </w:tabs>
        ''')
        pPr.append(tabs)
        
        b_tag = "<w:b/>" if level == 1 else ""
        color_val = "003366" if level == 1 else "000000"
        sz_val = "26" if level == 1 else ("24" if level == 2 else "22")
        
        hyperlink = parse_xml(f'''
            <w:hyperlink {nsdecls("w")} w:anchor="{bookmark}">
                <w:r>
                    <w:rPr>
                        <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                        <w:color w:val="{color_val}"/>
                        <w:sz w:val="{sz_val}"/>
                        {b_tag}
                    </w:rPr>
                    <w:t>{title}</w:t>
                </w:r>
                <w:r>
                    <w:rPr>
                        <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    </w:rPr>
                    <w:tab/>
                </w:r>
                <w:r>
                    <w:rPr>
                        <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                        <w:color w:val="{color_val}"/>
                        <w:sz w:val="{sz_val}"/>
                        {b_tag}
                    </w:rPr>
                    <w:t>{page}</w:t>
                </w:r>
            </w:hyperlink>
        ''')
        p._p.append(hyperlink)

    # Pre-populated Table of Contents items
    toc_items = [
        ("LỜI CẢM ƠN", "sec_thanks", 1, 2),
        ("DANH MỤC THUẬT NGỮ VÀ TỪ VIẾT TẮT", "sec_abbrev", 1, 4),
        ("DANH MỤC HÌNH ẢNH", "sec_figures", 1, 5),
        ("DANH MỤC BẢNG BIỂU", "sec_tables", 1, 6),
        ("LỜI NÓI ĐẦU", "sec_preface", 1, 7),
        ("CHƯƠNG 1: TỔNG QUAN ĐỀ TÀI", "sec_ch1", 1, 8),
        ("1.1. Mục tiêu của đề tài", "sec_1_1", 2, 8),
        ("1.2. Các bước hoàn thành đề tài", "sec_1_2", 2, 8),
        ("CHƯƠNG 2: CƠ SỞ LÝ THUYẾT", "sec_ch2", 1, 10),
        ("2.1. Hệ quản trị cơ sở dữ liệu MySQL", "sec_2_1", 2, 10),
        ("2.1.1. Khái niệm", "sec_2_1_1", 3, 10),
        ("2.1.2. Đặc điểm", "sec_2_1_2", 3, 10),
        ("2.1.3. Lợi ích khi sử dụng MySQL", "sec_2_1_3", 3, 10),
        ("2.2. Tổng quan về ngôn ngữ Java", "sec_2_2", 2, 12),
        ("2.2.1. Khái niệm", "sec_2_2_1", 3, 12),
        ("2.2.2. Đặc điểm của Java", "sec_2_2_2", 3, 12),
        ("2.2.3. Các platform của Java", "sec_2_2_3", 3, 13),
        ("2.2.4. Tiêu chuẩn của một môi trường Java điển hình", "sec_2_2_4", 3, 13),
        ("2.3. Tổng quan về Framework Spring Boot", "sec_2_3", 2, 14),
        ("2.3.1. Khái niệm và nguyên lý hoạt động", "sec_2_3_1", 3, 14),
        ("2.3.2. Ưu điểm nổi bật của Spring Boot trong dự án", "sec_2_3_2", 3, 14),
        ("2.4. Cơ chế bảo mật Spring Security và JSON Web Token (JWT)", "sec_2_4", 2, 15),
        ("2.4.1. Kiến trúc bảo mật Spring Security", "sec_2_4_1", 3, 15),
        ("2.4.2. Cơ chế xác thực phi trạng thái bằng JWT", "sec_2_4_2", 3, 16),
        ("2.4.3. Xác thực người dùng qua mạng xã hội OAuth2", "sec_2_4_3", 3, 16),
        ("2.5. Giao tiếp thời gian thực với WebSocket và STOMP", "sec_2_5", 2, 17),
        ("2.5.1. Khái niệm và cơ chế hoạt động của WebSocket", "sec_2_5_1", 3, 17),
        ("2.5.2. Giao thức STOMP và thư viện SockJS", "sec_2_5_2", 3, 17),
        ("2.6. Tổng quan về ReactJS", "sec_2_6", 2, 18),
        ("2.6.1. Khái niệm và các tính năng cốt lõi", "sec_2_6_1", 3, 18),
        ("2.6.2. Ưu và nhược điểm của ReactJS", "sec_2_6_2", 3, 18),
        ("2.7. Tổng quan về StarUML", "sec_2_7", 2, 19),
        ("2.7.1. Khái niệm và vai trò", "sec_2_7_1", 3, 19),
        ("2.7.2. Thực thể, thuộc tính và liên kết", "sec_2_7_2", 3, 19),
        ("CHƯƠNG 3: PHÂN TÍCH THIẾT KẾ HỆ THỐNG", "sec_ch3", 1, 21),
        ("3.1. Các bước hoàn thành dự án", "sec_3_1", 2, 21),
        ("3.2. Phân tích thiết kế hệ thống", "sec_3_2", 2, 22),
        ("3.2.1. Tác nhân và ca sử dụng", "sec_3_2_1", 3, 22),
        ("3.2.2. Danh sách Use Case", "sec_3_2_2", 3, 22),
        ("3.2.3. Đặc tả chi tiết các ca sử dụng", "sec_3_2_3", 3, 23),
        ("3.3. Biểu đồ ca sử dụng (Use Case Diagram)", "sec_3_3", 2, 25),
        ("3.4. Biểu đồ trình tự (Sequence Diagram)", "sec_3_4", 2, 27),
        ("3.4.1. Quản lý thuê phòng", "sec_3_4_1", 3, 27),
        ("3.4.2. Quản lý hợp đồng", "sec_3_4_2", 3, 28),
        ("3.4.3. Quản lý phòng, bài đăng", "sec_3_4_3", 3, 29),
        ("3.4.4. Quản lý bảo trì", "sec_3_4_4", 3, 30),
        ("3.4.5. Quản lý yêu cầu", "sec_3_4_5", 3, 31),
        ("3.4.6. Quản lý tài khoản", "sec_3_4_6", 3, 32),
        ("3.4.7. Chức năng gỡ và phê duyệt bài đăng", "sec_3_4_7", 3, 33),
        ("3.5. Biểu đồ trạng thái (State Diagram)", "sec_3_5", 2, 34),
        ("3.5.1. Trạng thái Tài khoản", "sec_3_5_1", 3, 34),
        ("3.5.2. Trạng thái Phòng", "sec_3_5_2", 3, 34),
        ("3.5.3. Trạng thái Bài đăng của phòng", "sec_3_5_3", 3, 35),
        ("3.5.4. Trạng thái Xử lý yêu cầu", "sec_3_5_4", 3, 35),
        ("3.6. Biểu đồ hoạt động (Activity Diagram)", "sec_3_6", 2, 36),
        ("3.6.1. Chức năng đăng nhập", "sec_3_6_1", 3, 36),
        ("3.6.2. Chức năng đăng ký", "sec_3_6_2", 3, 37),
        ("3.6.3. Chức năng lập hợp đồng thuê phòng", "sec_3_6_3", 3, 38),
        ("3.6.4. Chức năng lập phiếu bảo trì", "sec_3_6_4", 3, 39),
        ("3.6.5. Chức năng lập phiếu yêu cầu", "sec_3_6_5", 3, 40),
        ("3.6.6. Chức năng thêm phòng, bài đăng", "sec_3_6_6", 3, 41),
        ("3.6.7. Chức năng phê duyệt phòng", "sec_3_6_7", 3, 42),
        ("3.6.8. Chức năng gỡ phòng, bài đăng", "sec_3_6_8", 3, 43),
        ("3.6.9. Chức năng phân quyền", "sec_3_6_9", 3, 44),
        ("3.6.10. Chức năng bình luận", "sec_3_6_10", 3, 45),
        ("3.6.11. Chức năng đánh giá", "sec_3_6_11", 3, 46),
        ("3.7. Biểu đồ lớp chi tiết (Class Diagram)", "sec_3_7", 2, 47),
        ("3.8. Biểu đồ thành phần (Component Diagram)", "sec_3_8", 2, 48),
        ("3.8.1. Người quản trị", "sec_3_8_1", 3, 48),
        ("3.8.2. Người cho thuê", "sec_3_8_2", 3, 49),
        ("3.9. Thiết kế cơ sở dữ liệu", "sec_3_9", 2, 50),
        ("3.9.1. Cơ sở dữ liệu và mối quan hệ (ERD)", "sec_3_9_1", 3, 50),
        ("3.9.2. Mô tả chi tiết các bảng dữ liệu", "sec_3_9_2", 3, 50),
        ("CHƯƠNG 4: KẾT QUẢ CÀI ĐẶT CHƯƠNG TRÌNH", "sec_ch4", 1, 60),
        ("4.1. Chức năng dành cho người thuê (Khách hàng)", "sec_4_1", 2, 60),
        ("4.2. Chức năng dành cho người cho thuê (Chủ trọ)", "sec_4_2", 2, 62),
        ("4.3. Chức năng dành cho quản trị viên (Admin)", "sec_4_3", 2, 64),
        ("KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", "sec_conclusion", 1, 66),
        ("TÀI LIỆU THAM KHẢO", "sec_ref", 1, 68),
    ]

    for title, bm, lvl, pnum in toc_items:
        add_toc_row(title, bm, lvl, pnum)

    # End TOC Field
    p_toc_end = doc.add_paragraph()
    r4 = p_toc_end.add_run()
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    r4._r.append(fld3)

    doc.add_page_break()

    # =========================================================================
    # 4. DANH MỤC THUẬT NGỮ VÀ TỪ VIẾT TẮT
    # =========================================================================
    add_heading("DANH MỤC THUẬT NGỮ VÀ TỪ VIẾT TẮT", level=1, bookmark="sec_abbrev")
    
    abbrevs = [
        ("API", "Application Programming Interface", "Giao diện lập trình ứng dụng"),
        ("JWT", "JSON Web Token", "Chuẩn mở định dạng an toàn cho việc xác thực và truyền tải dữ liệu"),
        ("OAuth2", "Open Authorization 2.0", "Giao thức ủy quyền tiêu chuẩn bảo mật cho phép đăng nhập qua bên thứ 3"),
        ("STOMP", "Simple Text Oriented Messaging Protocol", "Giao thức định hướng văn bản đơn giản trao đổi tin nhắn qua WebSocket"),
        ("WebSocket", "WebSocket Protocol", "Giao thức mạng cho phép truyền thông hai chiều thời gian thực qua một kết nối TCP"),
        ("ORM", "Object-Relational Mapping", "Kỹ thuật ánh xạ giữa các bảng trong CSDL và các đối tượng mã nguồn"),
        ("JPA", "Java Persistence API", "Giao diện đặc tả quản lý dữ liệu quan hệ trong ứng dụng Java"),
        ("DOM", "Document Object Model", "Mô hình các đối tượng trong tài liệu HTML biểu diễn cấu trúc cây trang web"),
        ("SPA", "Single Page Application", "Ứng dụng web tải một trang duy nhất và cập nhật nội dung động mà không tải lại trang"),
        ("CRUD", "Create, Read, Update, Delete", "Bốn thao tác cơ bản với cơ sở dữ liệu: Thêm, Xem, Sửa, Xóa"),
        ("GUI", "Graphical User Interface", "Giao diện đồ họa người dùng"),
        ("HTML", "HyperText Markup Language", "Ngôn ngữ đánh dấu siêu văn bản định hình nội dung web"),
        ("IDE", "Integrated Development Environment", "Môi trường phát triển tích hợp dành cho lập trình viên"),
        ("TCP", "Transmission Control Protocol", "Giao thức điều khiển truyền vận kết nối tin cậy"),
        ("SSL", "Secure Sockets Layer", "Giao thức mã hóa an toàn truyền tải dữ liệu trên Internet"),
        ("RDBMS", "Relational Database Management System", "Hệ quản trị cơ sở dữ liệu quan hệ"),
        ("RBAC", "Role-Based Access Control", "Cơ chế kiểm soát truy cập phân quyền dựa trên vai trò người dùng")
    ]
    add_styled_table(["Thuật ngữ / Từ viết tắt", "Cụm từ đầy đủ", "Ý nghĩa chuyên môn"], abbrevs, [1.5, 2.5, 2.5])

    doc.add_page_break()

    # =========================================================================
    # 5. DANH MỤC HÌNH ẢNH & BẢNG BIỂU
    # =========================================================================
    add_heading("DANH MỤC HÌNH ẢNH", level=1, bookmark="sec_figures")
    figures_list = [
        ("Hình 2.1", "Lợi ích khi sử dụng MySQL"),
        ("Hình 2.2", "Platform cơ bản của Java"),
        ("Hình 2.3", "Môi trường Java điển hình"),
        ("Hình 2.4", "Mô hình kiến trúc Spring Boot"),
        ("Hình 2.5", "Cơ chế xác thực qua JWT"),
        ("Hình 3.1", "Biểu đồ ca sử dụng của người cho thuê"),
        ("Hình 3.2", "Biểu đồ ca sử dụng của quản trị viên"),
        ("Hình 3.3", "Biểu đồ phân quyền của quản trị viên"),
        ("Hình 3.4", "Biểu đồ ca sử dụng thuê phòng"),
        ("Hình 3.5", "Biểu đồ ca sử dụng đánh giá và bình luận"),
        ("Hình 3.6", "Biểu đồ trình tự thuê phòng"),
        ("Hình 3.7", "Biểu đồ trình tự hợp đồng"),
        ("Hình 3.8", "Biểu đồ trình tự phòng, bài đăng"),
        ("Hình 3.9", "Biểu đồ trình tự phiếu bảo trì"),
        ("Hình 3.10", "Biểu đồ trình tự phiếu yêu cầu"),
        ("Hình 3.11", "Biểu đồ trình tự quản lý tài khoản"),
        ("Hình 3.12", "Biểu đồ trình tự phê duyệt và gỡ phòng, bài đăng"),
        ("Hình 3.13", "Biểu đồ trạng thái của tài khoản"),
        ("Hình 3.14", "Biểu đồ trạng thái của phòng"),
        ("Hình 3.15", "Biểu đồ trạng thái của bài đăng phòng"),
        ("Hình 3.16", "Biểu đồ trạng thái của xử lý yêu cầu"),
        ("Hình 3.17", "Biểu đồ hoạt động của chức năng đăng nhập"),
        ("Hình 3.18", "Biểu đồ hoạt động của chức năng đăng ký"),
        ("Hình 3.19", "Biểu đồ hoạt động của chức năng lập hợp đồng thuê phòng"),
        ("Hình 3.20", "Biểu đồ hoạt động của chức năng lập phiếu bảo trì"),
        ("Hình 3.21", "Biểu đồ hoạt động của chức năng lập phiếu yêu cầu"),
        ("Hình 3.22", "Biểu đồ hoạt động của chức năng thêm phòng, bài đăng"),
        ("Hình 3.23", "Biểu đồ hoạt động của chức năng phê duyệt phòng"),
        ("Hình 3.24", "Biểu đồ hoạt động của chức năng gỡ bài đăng"),
        ("Hình 3.25", "Biểu đồ hoạt động của chức năng phân quyền"),
        ("Hình 3.26", "Biểu đồ hoạt động của chức năng bình luận"),
        ("Hình 3.27", "Biểu đồ hoạt động của chức năng đánh giá"),
        ("Hình 3.28", "Biểu đồ lớp chi tiết"),
        ("Hình 3.29", "Biểu đồ thành phần người quản trị"),
        ("Hình 3.30", "Biểu đồ thành phần người cho thuê"),
        ("Hình 3.31", "Cơ sở dữ liệu và mối quan hệ (ERD)"),
        ("Hình 4.1", "Giao diện Trang chủ và Danh sách phòng trọ"),
        ("Hình 4.2", "Giao diện Chi tiết phòng trọ và Bản đồ Google Maps"),
        ("Hình 4.3", "Giao diện Gửi yêu cầu đặt thuê phòng"),
        ("Hình 4.4", "Giao diện Đánh giá, bình luận phòng trọ"),
        ("Hình 4.5", "Giao diện Nhắn tin trực tuyến (Chat Realtime)"),
        ("Hình 4.6", "Giao diện Quản lý danh sách phòng đã thuê"),
        ("Hình 4.7", "Giao diện Dashboard thống kê doanh thu và chi phí theo tháng"),
        ("Hình 4.8", "Giao diện Quản lý danh sách phòng và Thêm mới phòng"),
        ("Hình 4.9", "Giao diện Quản lý và Lập hợp đồng thuê phòng"),
        ("Hình 4.10", "Giao diện Quản lý chỉ số Điện - Nước và tính tiền hàng tháng"),
        ("Hình 4.11", "Giao diện Quản lý yêu cầu từ khách thuê và xuất hóa đơn"),
        ("Hình 4.12", "Giao diện Quản lý bảo trì và chi phí bảo trì phòng"),
        ("Hình 4.13", "Giao diện Dashboard quản trị hệ thống Admin"),
        ("Hình 4.14", "Giao diện Phê duyệt và Gỡ bỏ bài đăng phòng trọ"),
        ("Hình 4.15", "Giao diện Quản lý tài khoản và Phân quyền người dùng"),
        ("Hình 4.16", "Giao diện Gửi email thông báo tới người dùng")
    ]
    add_styled_table(["Mã số hình", "Tên hình ảnh minh họa / Sơ đồ"], figures_list, [1.5, 5.0])

    add_heading("DANH MỤC BẢNG BIỂU", level=1, bookmark="sec_tables")
    tables_list = [
        ("Bảng 3.1", "Các bước hoàn thành dự án"),
        ("Bảng 3.2", "Bảng User (Người dùng)"),
        ("Bảng 3.3", "Bảng Role (Quyền người dùng)"),
        ("Bảng 3.4", "Bảng Room (Phòng trọ)"),
        ("Bảng 3.5", "Bảng Rate (Đánh giá sao)"),
        ("Bảng 3.6", "Bảng Follow (Theo dõi chủ trọ)"),
        ("Bảng 3.7", "Bảng BlogStore (Lưu bài đăng)"),
        ("Bảng 3.8", "Bảng Message (Cuộc hội thoại)"),
        ("Bảng 3.9", "Bảng MessageChat (Chi tiết tin nhắn)"),
        ("Bảng 3.10", "Bảng Asset (Tài sản phòng trọ)"),
        ("Bảng 3.11", "Bảng Comment (Bình luận)"),
        ("Bảng 3.12", "Bảng Category (Danh mục phòng trọ)"),
        ("Bảng 3.13", "Bảng Location (Khu vực / Thành phố)"),
        ("Bảng 3.14", "Bảng Maintenance (Phiếu bảo trì)"),
        ("Bảng 3.15", "Bảng Request (Phiếu yêu cầu khách thuê)"),
        ("Bảng 3.16", "Bảng Contract (Hợp đồng thuê phòng)"),
        ("Bảng 3.17", "Bảng RoomMedia (Hình ảnh phòng trọ)"),
        ("Bảng 3.18", "Bảng ElectricAndWater (Quản lý chỉ số Điện & Nước hàng tháng)")
    ]
    add_styled_table(["Mã số bảng", "Tên bảng dữ liệu / Mô tả bảng"], tables_list, [1.5, 5.0])

    doc.add_page_break()

    # =========================================================================
    # 6. LỜI NÓI ĐẦU
    # =========================================================================
    add_heading("LỜI NÓI ĐẦU", level=1, bookmark="sec_preface")
    add_p("Khi xã hội ngày càng phát triển và quá trình đô thị hóa diễn ra mạnh mẽ, mật độ dân số tập trung tại các thành phố lớn ngày một tăng cao. Đặc biệt với các đối tượng như sinh viên các trường đại học, cao đẳng, sinh viên mới ra trường, các gia đình trẻ mới kết hôn, công nhân tại các khu công nghiệp hay người lao động ngoại tỉnh thì nhu cầu về nhà ở là hết sức cấp thiết. Khả năng tài chính ban đầu thường chưa đủ để sở hữu một căn nhà hay căn hộ chung cư cao cấp, do đó phân khúc nhà trọ và căn hộ mini cho thuê với mức giá hợp lý luôn là sự lựa chọn hàng đầu.")
    add_p("Thế nhưng, việc xây dựng và vận hành nhà trọ lại đặt ra rất nhiều thách thức trong công tác quản lý. Khi quy mô số lượng phòng trọ tăng lên và lượng khách trọ đông đảo, các công việc như: quản lý thông tin khách thuê, lập và theo dõi thời hạn hợp đồng, ghi nhận và tính toán số tiền điện nước biến động từng tháng, quản lý trang thiết bị tài sản trong từng phòng, xử lý các sự cố bảo trì và đặc biệt là báo cáo thống kê doanh thu - chi phí cần được thực hiện chặt chẽ, chính xác và kịp thời.")
    add_p("Phương thức quản lý truyền thống bằng sổ sách giấy tờ hoặc file Excel rời rạc bộc lộ nhiều điểm hạn chế: tốn nhiều công sức ghi chép, dễ xảy ra sai sót nhầm lẫn trong quá trình tính toán tiền dịch vụ, dữ liệu dễ bị mất mát và rất khó khăn trong việc tra cứu lại lịch sử hợp đồng hay thống kê dòng tiền.")
    add_p("Xuất phát từ thực tiễn trên, việc xây dựng một hệ thống website toàn diện phục vụ cả người thuê và người cho thuê là một yêu cầu cấp bách. Hệ thống không chỉ giúp người thuê dễ dàng tìm kiếm phòng trọ đúng nhu cầu, định vị chính xác vị trí trên bản đồ và liên hệ nhanh chóng, mà còn giúp các chủ nhà trọ số hóa toàn bộ quy trình vận hành: quản lý tin đăng, hợp đồng, tự động tính toán điện nước và theo dõi biểu đồ tài chính trực quan. Đề tài 'Nghiên cứu xây dựng hệ thống website quản lý và cho thuê nhà trọ' được thực hiện nhằm giải quyết trọn vẹn các bài toán thực tiễn nêu trên.")

    doc.add_page_break()

    # =========================================================================
    # CHƯƠNG 1: TỔNG QUAN ĐỀ TÀI
    # =========================================================================
    add_heading("CHƯƠNG 1: TỔNG QUAN ĐỀ TÀI", level=1, bookmark="sec_ch1")
    
    add_heading("1.1. Mục tiêu của đề tài", level=2, bookmark="sec_1_1")
    add_p("Đề tài tập trung giải quyết các mục tiêu trọng tâm sau đây:")
    add_p("1. Tối ưu hóa quy trình quản lý và cho thuê nhà trọ: Xây dựng một nền tảng website tập trung giúp kết nối hiệu quả giữa chủ nhà trọ và khách thuê; chuẩn hóa quy trình đăng tin, duyệt bài, lập hợp đồng thuê và bàn giao phòng.")
    add_p("2. Tự động hóa công tác tính toán và quản trị chi phí: Xây dựng module quản lý chỉ số điện nước theo tháng, hỗ trợ ghi nhận số điện cũ/mới, số khối nước cũ/mới, đơn giá và tự động tính ra tổng tiền; hỗ trợ xuất các phiếu thu và hóa đơn thanh lý trả phòng.")
    add_p("3. Xây dựng giao diện người dùng thân thiện, trực quan: Ứng dụng công nghệ Single Page Application (ReactJS) tạo ra trải nghiệm mượt mà, phân cấp giao diện thông minh cho từng vai trò (Khách thuê, Chủ trọ, Quản trị viên).")
    add_p("4. Nâng cao tính năng tương tác và định vị: Tích hợp Google Maps API giúp người dùng dễ dàng định vị vị trí phòng trọ; tích hợp hệ thống trò chuyện trực tuyến (Chat Realtime) thời gian thực sử dụng WebSocket và STOMP giữa khách thuê và chủ nhà.")
    add_p("5. Tăng cường tính bảo mật và phân quyền chặt chẽ: Áp dụng các chuẩn bảo mật công nghiệp gồm Spring Security, JWT (JSON Web Token) và hỗ trợ xác thực liên kết mạng xã hội (Google, Facebook OAuth2).")
    add_p("6. Thống kê và hỗ trợ ra quyết định: Cung cấp biểu đồ trực quan (Chart.js) theo dõi doanh thu theo tháng, cơ cấu chi phí phụ trợ và tỷ lệ lấp đầy phòng trọ.")

    add_heading("1.2. Các bước hoàn thành đề tài", level=2, bookmark="sec_1_2")
    add_p("Để hoàn thành đồ án một cách có hệ thống, đề tài được triển khai qua các bước cụ thể:")
    add_p("1. Phân tích yêu cầu và thiết kế cơ sở dữ liệu: Khảo sát thực tế nhu cầu của khách thuê và chủ trọ; thiết kế cơ sở dữ liệu quan hệ trên MySQL lưu trữ đầy đủ thực thể (phòng trọ, hợp đồng, điện nước, hóa đơn, tài khoản, đánh giá, tin nhắn).")
    add_p("2. Xây dựng backend với Java Spring Boot: Phát triển hệ thống RESTful API cung cấp dữ liệu; xây dựng cơ chế xác thực JWT & OAuth2; xử lý logic nghiệp vụ quản lý phòng, hợp đồng, tính toán hóa đơn và thống kê.")
    add_p("3. Xây dựng frontend với ReactJS: Thiết kế giao diện người dùng trực quan, responsive; giao tiếp với API backend để truy xuất và cập nhật dữ liệu; tối ưu hóa trải nghiệm tìm kiếm, lọc và xem bản đồ.")
    add_p("4. Tích hợp Real-time và Tiện ích ngoại vi: Xây dựng module Chat thời gian thực với WebSocket và STOMP; tích hợp dịch vụ gửi email tự động (JavaMailSender) và xuất báo cáo/hóa đơn (Apache POI).")
    add_p("5. Kiểm thử và đánh giá hiệu năng: Thực hiện các kịch bản kiểm thử chức năng, tính toán dòng tiền, độ bảo mật và tính đáp ứng của hệ thống.")
    add_p("6. Triển khai thử nghiệm và đánh giá: Triển khai ứng dụng trên môi trường thử nghiệm và viết tài liệu hướng dẫn.")

    doc.add_page_break()

    # =========================================================================
    # CHƯƠNG 2: CƠ SỞ LÝ THUYẾT
    # =========================================================================
    add_heading("CHƯƠNG 2: CƠ SỞ LÝ THUYẾT", level=1, bookmark="sec_ch2")

    add_heading("2.1. Hệ quản trị cơ sở dữ liệu MySQL", level=2, bookmark="sec_2_1")
    add_heading("2.1.1. Khái niệm", level=3, bookmark="sec_2_1_1")
    add_p("Hệ quản trị cơ sở dữ liệu MySQL được hiểu như là chương trình dùng để quản lý hệ thống cơ sở dữ liệu, trong đó, cơ sở dữ liệu là một hệ thống lưu trữ thông tin được sắp xếp rõ ràng, phân lớp ngăn nắp. Nó giúp người dùng có thể truy cập dữ liệu một cách thuận lợi và nhanh chóng nhất. Vì hỗ trợ đa số các ngôn ngữ lập trình nên MySQL chính là hệ quản trị cơ sở dữ liệu tự do nguồn mở phổ biến nhất trên thế giới. Hiện MySQL đang được các nhà phát triển rất ưa chuộng trong quá trình phát triển ứng dụng.")
    add_p("MySQL là cơ sở dữ liệu tốc độ cao, ổn định hoạt động trên nhiều hệ điều hành, cung cấp một hệ thống lớn các hàm tiện ích rất mạnh. Đặc biệt, hệ quản trị cơ sở dữ liệu MySQL hoàn toàn miễn phí nên người dùng có thể thoải mái tải về từ trang chủ. Nó có rất nhiều phiên bản cho các hệ điều hành khác nhau như Windows, Linux, macOS.")

    add_heading("2.1.2. Đặc điểm", level=3, bookmark="sec_2_1_2")
    add_p("Hệ quản trị cơ sở dữ liệu MySQL chính là một trong những phần mềm quản trị CSDL dạng server-based, hoạt động theo mô hình máy khách - máy chủ.")
    add_p("MySQL quản lý dữ liệu thông qua các bảng quan hệ chứa dữ liệu riêng biệt, cho phép thiết lập các khóa chính, khóa ngoại để đảm bảo tính toàn vẹn dữ liệu tham chiếu.")
    add_p("MySQL có cơ chế phân quyền người dùng chặt chẽ, mỗi tài khoản truy cập đều được cấp phát quyền hạn cụ thể trên từng database, bảng hoặc câu lệnh (SELECT, INSERT, UPDATE, DELETE).")
    add_p("MySQL tương thích rất tốt với môi trường Java thông qua giao thức JDBC (Java Database Connectivity), giúp hệ thống hoạt động mạnh mẽ, ổn định.")

    add_heading("2.1.3. Lợi ích khi sử dụng MySQL", level=3, bookmark="sec_2_1_3")
    add_image_placeholder("Hình 2.1", "Lợi ích khi sử dụng MySQL")
    add_p("Ra mắt người dùng hoàn toàn miễn phí, hệ quản trị cơ sở dữ liệu MySQL mang đến rất nhiều lợi ích thiết thực, được ứng dụng rộng rãi từ các trang web cá nhân đến các hệ thống doanh nghiệp lớn:")
    add_p("- Hiệu năng sử dụng cao: Với kiến trúc Storage Engine linh hoạt (đặc biệt là InnoDB), MySQL mang lại hiệu năng truy vấn vượt trội, xử lý đồng thời hàng nghìn kết nối với độ trễ thấp.")
    add_p("- Hỗ trợ giao dịch mạnh mẽ (ACID): Hỗ trợ đầy đủ các tiêu chuẩn Atomicity, Consistency, Isolation, Durability, đảm bảo an toàn tuyệt đối cho các giao dịch lập hợp đồng và thanh toán tiền phòng.")
    add_p("- Khóa mức dòng (Row-level Locking): Hạn chế xung đột tài nguyên khi nhiều người dùng cùng thao tác cập nhật dữ liệu đồng thời.")
    add_p("- Tính bảo mật và kết nối cao: Tích hợp các lớp bảo mật mạnh mẽ, hỗ trợ kết nối an toàn qua giao thức SSL/TLS và mã hóa mật khẩu người dùng.")
    add_p("- Chi phí sở hữu thấp: Là phần mềm mã nguồn mở hoàn toàn miễn phí, tiết kiệm tối đa ngân sách đầu tư hạ tầng phần mềm cho dự án.")

    add_heading("2.2. Tổng quan về ngôn ngữ Java", level=2, bookmark="sec_2_2")
    add_heading("2.2.1. Khái niệm", level=3, bookmark="sec_2_2_1")
    add_p("Java là ngôn ngữ lập trình bậc cao, hướng đối tượng hoàn toàn, được Sun Microsystems phát triển vào năm 1995 (nay thuộc sở hữu của Oracle). Java kế thừa cú pháp trong sáng từ C/C++ nhưng loại bỏ đi các cấu trúc phức tạp và tiềm ẩn nguy cơ lỗi như con trỏ trực tiếp và đa kế thừa lớp.")

    add_heading("2.2.2. Đặc điểm của Java", level=3, bookmark="sec_2_2_2")
    add_p("Tiêu chí hàng đầu của Ngôn ngữ Lập trình Java là 'Write Once, Run Anywhere' (Viết một lần, chạy mọi nơi). Java biên dịch mã nguồn thành bytecode độc lập với phần cứng, cho phép chương trình chạy được trên mọi nền tảng hệ điều hành có cài đặt Máy ảo Java (JVM).")
    add_p("- Đơn giản và quen thuộc: Cú pháp gần gũi với C/C++ nhưng lược bỏ các cơ chế dễ gây lỗi bộ nhớ.")
    add_p("- Hướng đối tượng chặt chẽ: Mọi thành phần trong Java đều được tổ chức dưới dạng lớp (Class) và đối tượng (Object), hỗ trợ đầy đủ 4 tính chất: Đóng gói, Kế thừa, Đa hình và Trừu tượng.")
    add_p("- An toàn và mạnh mẽ: Tích hợp cơ chế thu gom rác tự động (Garbage Collection), tự động giải phóng vùng nhớ khi đối tượng không còn được tham chiếu.")
    add_p("- Đa nhiệm (Multithreading): Cho phép thực hiện nhiều luồng công việc đồng thời, tối ưu hóa việc xử lý các yêu cầu từ nhiều người dùng trên máy chủ web.")

    add_heading("2.2.3. Các platform của Java", level=3, bookmark="sec_2_2_3")
    add_image_placeholder("Hình 2.2", "Platform cơ bản của Java")
    add_p("Nền tảng Java bao gồm 3 thành phần trụ cột:")
    add_p("- Java Virtual Machine (JVM): Máy ảo thực thi mã bytecode (.class).")
    add_p("- Java Development Kit (JDK): Bộ công cụ dành cho lập trình viên gồm trình biên dịch (javac), trình thông dịch (java) và các công cụ phát triển.")
    add_p("- Java Runtime Environment (JRE): Môi trường cung cấp thư viện chuẩn và JVM để thực thi ứng dụng.")

    add_heading("2.2.4. Tiêu chuẩn của một môi trường Java điển hình", level=3, bookmark="sec_2_2_4")
    add_image_placeholder("Hình 2.3", "Môi trường Java điển hình")
    add_p("Chương trình Java trải qua 5 giai đoạn chính:")
    add_p("1. Soạn thảo (Editor): Lập trình viên viết mã nguồn lưu trong tệp có đuôi mở rộng là .java.")
    add_p("2. Biên dịch (Compiler): Trình biên dịch javac chuyển mã nguồn thành tệp chứa mã bytecode (.class).")
    add_p("3. Nạp lớp (Class Loader): Nạp các tệp .class vào bộ nhớ của máy ảo JVM.")
    add_p("4. Kiểm tra mã Bytecode (Bytecode Verifier): Kiểm tra tính hợp lệ và đảm bảo mã không vi phạm các quy tắc bảo mật của Java.")
    add_p("5. Thực thi (Interpreter / JIT Compiler): Máy ảo chuyển đổi bytecode thành mã máy để CPU xử lý.")

    add_heading("2.3. Tổng quan về Framework Spring Boot", level=2, bookmark="sec_2_3")
    add_heading("2.3.1. Khái niệm và nguyên lý hoạt động", level=3, bookmark="sec_2_3_1")
    add_p("Spring Boot là framework mã nguồn mở được phát triển bởi Pivotal Software (nay thuộc VMware), xây dựng trên nền tảng Spring Framework kinh điển nhằm đơn giản hóa việc khởi tạo, cấu hình và triển khai các ứng dụng web độc lập cấp doanh nghiệp (Enterprise Applications).")
    add_p("Nguyên lý hoạt động của Spring Boot dựa trên hai trụ cột chính: Đảo ngược điều khiển (Inversion of Control - IoC) và Tiêm phụ thuộc (Dependency Injection - DI). Thông qua các annotation như @Service, @Repository, @Controller, @Autowired, Spring Boot tự động quản lý vòng đời của các bean và khởi tạo các đối tượng mà không cần cấu hình file XML phức tạp.")
    add_image_placeholder("Hình 2.4", "Mô hình kiến trúc Spring Boot")

    add_heading("2.3.2. Ưu điểm nổi bật của Spring Boot trong dự án", level=3, bookmark="sec_2_3_2")
    add_p("- Máy chủ nhúng (Embedded Server): Tích hợp sẵn máy chủ Tomcat, giúp ứng dụng có thể đóng gói thành một tệp JAR duy nhất và thực thi trực tiếp bằng lệnh java -jar.")
    add_p("- Tự động cấu hình (Auto-Configuration): Tự động phát hiện các thư viện trong classpath và cấu hình các thành phần tương ứng một cách thông minh.")
    add_p("- Spring Data JPA: Cung cấp abstraction layer mạnh mẽ bên trên Hibernate ORM. Lập trình viên chỉ cần khai báo interface kế thừa JpaRepository là có ngay các hàm CRUD chuẩn và phân trang mà không cần viết câu lệnh SQL thủ công.")
    add_p("- Starter Dependencies: Cung cấp các gói dependency tổng hợp (như spring-boot-starter-web, spring-boot-starter-security, spring-boot-starter-data-jpa) giúp quản lý phiên bản thư viện đồng bộ và nhất quán.")

    add_heading("2.4. Cơ chế bảo mật Spring Security và JSON Web Token (JWT)", level=2, bookmark="sec_2_4")
    add_heading("2.4.1. Kiến trúc bảo mật Spring Security", level=3, bookmark="sec_2_4_1")
    add_p("Spring Security là chuẩn bảo mật mạnh mẽ và toàn diện dành cho các ứng dụng Java Spring. Hệ thống hoạt động dựa trên cơ chế chuỗi các bộ lọc bảo mật (Security Filter Chain) đánh chặn tất cả các HTTP request gửi đến máy chủ trước khi chúng chạm tới các Controller xử lý nghiệp vụ.")
    add_p("Trong dự án, Spring Security đảm bảo việc phân quyền chặt chẽ theo vai trò (Role-Based Access Control): Quản trị viên (ROLE_ADMIN), Chủ trọ (ROLE_RENTALER) và Khách thuê (ROLE_USER).")

    add_heading("2.4.2. Cơ chế xác thực phi trạng thái bằng JWT", level=3, bookmark="sec_2_4_2")
    add_image_placeholder("Hình 2.5", "Cơ chế xác thực qua JWT")
    add_p("Trong kiến trúc RESTful API hiện đại giữa ReactJS và Spring Boot, việc lưu trữ phiên đăng nhập bằng Session trên server gây tốn bộ nhớ và khó mở rộng. Do đó, hệ thống sử dụng cơ chế xác thực phi trạng thái (Stateless Authentication) qua JSON Web Token (JWT):")
    add_p("1. Đăng nhập: Người dùng gửi thông tin email và mật khẩu lên endpoint /auth/login.")
    add_p("2. Cấp phát Token: Máy chủ xác thực thông tin hợp lệ, tạo một chuỗi JWT được ký số bằng thuật toán HMAC SHA-256 chứa thông tin định danh (User ID, Role, thời gian hết hạn).")
    add_p("3. Xác thực request: Phía ReactJS lưu token vào LocalStorage. Trong mỗi request tiếp theo, client gửi token trong Header Authorization: Bearer <token>. Bộ lọc JwtAuthenticationFilter tại backend sẽ trích xuất token, giải mã và xác thực quyền truy cập.")

    add_heading("2.4.3. Xác thực người dùng qua mạng xã hội OAuth2", level=3, bookmark="sec_2_4_3")
    add_p("Hệ thống tích hợp thư viện spring-boot-starter-oauth2-client hỗ trợ người dùng đăng nhập nhanh chóng thông qua tài khoản Google hoặc Facebook. Sau khi người dùng xác thực thành công tại nhà cung cấp dịch vụ (Google/Facebook OAuth2 Provider), server sẽ tiếp nhận thông tin hồ sơ (Email, Name, Avatar) và tự động tạo hoặc đồng bộ tài khoản người dùng.")

    add_heading("2.5. Giao tiếp thời gian thực với WebSocket và STOMP", level=2, bookmark="sec_2_5")
    add_heading("2.5.1. Khái niệm và cơ chế hoạt động của WebSocket", level=3, bookmark="sec_2_5_1")
    add_p("WebSocket là giao thức mạng chuẩn hóa cho phép thiết lập kết nối truyền thông hai chiều toàn phần (Full-duplex) qua một kết nối TCP duy nhất giữa trình duyệt và máy chủ. Khác với giao thức HTTP truyền thống yêu cầu client phải liên tục gửi request thăm dò (polling), WebSocket giữ kết nối luôn mở, giúp server có thể chủ động đẩy dữ liệu tức thì xuống client ngay khi có sự kiện mới với độ trễ tính bằng mili-giây.")

    add_heading("2.5.2. Giao thức STOMP và thư viện SockJS", level=3, bookmark="sec_2_5_2")
    add_p("- STOMP (Simple Text Oriented Messaging Protocol): Là giao thức thông điệp định hướng văn bản chạy trên nền WebSocket. STOMP cung cấp mô hình Publish-Subscribe thông qua các kênh (Topic/Queue).")
    add_p("- SockJS: Là thư viện client-side cung cấp giải pháp dự phòng (fallback) tự động khi trình duyệt hoặc môi trường mạng của người dùng không hỗ trợ WebSocket thuần.")
    add_p("Trong hệ thống, module Nhắn tin trực tuyến sử dụng SockJS và STOMP kết nối tới endpoint /ws của Spring Boot, lắng nghe kênh /topic/messages giúp khách thuê và chủ phòng trọ có thể trò chuyện trực tiếp, trao đổi thông tin phòng tức thì mà không cần tải lại trang.")

    add_heading("2.6. Tổng quan về ReactJS", level=2, bookmark="sec_2_6")
    add_heading("2.6.1. Khái niệm và các tính năng cốt lõi", level=3, bookmark="sec_2_6_1")
    add_p("ReactJS là thư viện JavaScript mã nguồn mở được Facebook phát triển vào năm 2013, chuyên dùng để xây dựng giao diện người dùng cho ứng dụng web một trang (Single Page Application - SPA).")
    add_p("- JSX (JavaScript XML): Cú pháp mở rộng kết hợp HTML bên trong JavaScript, giúp việc mô tả cây giao diện trở nên trực quan, sinh động.")
    add_p("- Kiến trúc Component: Toàn bộ giao diện được phân tách thành các khối thành phần độc lập (Header, Sidebar, RoomCard, Modal), mỗi component tự quản lý trạng thái (state) và dữ liệu đầu vào (props), dễ dàng tái sử dụng.")
    add_p("- Virtual DOM (DOM ảo): React duy trì một bản sao cây DOM trên bộ nhớ RAM. Khi trạng thái component thay đổi, React so sánh Virtual DOM mới với Virtual DOM cũ (thuật toán Diffing) và chỉ cập nhật duy nhất những phần tử HTML thực sự thay đổi trên trình duyệt, mang lại tốc độ hiển thị cực kỳ mượt mà.")

    add_heading("2.6.2. Ưu và nhược điểm của ReactJS", level=3, bookmark="sec_2_6_2")
    add_p("- Ưu điểm: Tốc độ render giao diện cao, cộng đồng phát triển khổng lồ, hệ sinh thái thư viện phong phú (React Router v6, Material-UI, Ant Design, Axios, Chart.js), dễ bảo trì và mở rộng.")
    add_p("- Nhược điểm: Việc kết hợp JSX đòi hỏi thời gian làm quen ban đầu; thư viện phát triển nhanh yêu cầu lập trình viên phải liên tục cập nhật các chuẩn lập trình mới (Functional Component, React Hooks).")

    add_heading("2.7. Tổng quan về StarUML", level=2, bookmark="sec_2_7")
    add_heading("2.7.1. Khái niệm và vai trò", level=3, bookmark="sec_2_7_1")
    add_p("StarUML là công cụ mô hình hóa phần mềm đồ họa chuyên nghiệp hỗ trợ chuẩn ngôn ngữ mô hình hóa thống nhất UML (Unified Modeling Language). StarUML đóng vai trò cốt lõi trong giai đoạn phân tích thiết kế:")
    add_p("- Trực quan hóa cấu trúc và hành vi hệ thống qua các biểu đồ chuẩn.")
    add_p("- Kiểm tra tính hợp lệ và tính toàn vẹn của mô hình phân tích trước khi tiến hành viết code.")
    add_p("- Giúp các thành viên trong nhóm phát triển có cùng một góc nhìn thống nhất về nghiệp vụ.")

    add_heading("2.7.2. Thực thể, thuộc tính và liên kết", level=3, bookmark="sec_2_7_2")
    add_p("- Thực thể (Entity): Đại diện cho một tập hợp các đối tượng trong thế giới thực có cùng đặc tính cần quản lý (ví dụ: Người dùng, Phòng trọ, Hợp đồng, Hóa đơn điện nước).")
    add_p("- Thuộc tính (Attribute): Các thông tin chi tiết mô tả đặc trưng của thực thể (ví dụ: Mã phòng, Giá thuê, Địa chỉ, Số điện cũ, Số điện mới).")
    add_p("- Liên kết (Relationship): Biểu diễn mối quan hệ logic giữa các thực thể, gồm 3 kiểu liên kết cơ bản: 1 - 1 (Một - Một), 1 - N (Một - Nhiều) và N - N (Nhiều - Nhiều).")

    doc.add_page_break()

    # =========================================================================
    # CHƯƠNG 3: PHÂN TÍCH THIẾT KẾ HỆ THỐNG
    # =========================================================================
    add_heading("CHƯƠNG 3: PHÂN TÍCH THIẾT KẾ HỆ THỐNG", level=1, bookmark="sec_ch3")

    add_heading("3.1. Các bước hoàn thành dự án", level=2, bookmark="sec_3_1")
    add_p("Quá trình phân tích và phát triển hệ thống được tổ chức khoa học qua các giai đoạn chi tiết:")
    
    project_steps = [
        ("Khởi tạo dự án", "Khảo sát thực tế bài toán quản lý và cho thuê nhà trọ; nghiên cứu các quy định pháp lý về hợp đồng thuê nhà; đánh giá tính khả thi và phạm vi công nghệ."),
        ("Lập kế hoạch", "Xây dựng kế hoạch tổng quát; phân bổ các mốc hoàn thành; xác định rủi ro kỹ thuật và các giải pháp dự phòng."),
        ("Thu thập yêu cầu", "Khảo sát chi tiết nhu cầu nghiệp vụ của 3 đối tượng tác nhân: Khách tìm trọ, Chủ nhà trọ và Quản trị viên hệ thống."),
        ("Phân tích yêu cầu", "Phân loại các yêu cầu chức năng (CRUD phòng, hợp đồng, điện nước, chat, hóa đơn) và phi chức năng (bảo mật, hiệu năng, responsive)."),
        ("Mô hình hóa hệ thống", "Xây dựng hệ thống biểu đồ UML bằng StarUML: Use Case, Sequence, State, Activity, Class, Component và ERD CSDL."),
        ("Thiết kế CSDL & UI", "Thiết kế cơ sở dữ liệu quan hệ trên MySQL; thiết kế mẫu giao diện web trực quan, hiện đại."),
        ("Cài đặt hệ thống", "Lập trình Backend RESTful API với Spring Boot & Spring Security; lập trình Frontend SPA với ReactJS & Material-UI; tích hợp WebSocket."),
        ("Kiểm thử hệ thống", "Kiểm thử chức năng, kiểm thử luồng giao dịch tài chính, kiểm tra tính toàn vẹn dữ liệu và kiểm thử bảo mật."),
        ("Triển khai & Đánh giá", "Triển khai thử nghiệm hệ thống trên máy chủ; hiệu chỉnh các tham số cấu hình và hoàn thiện tài liệu báo cáo.")
    ]
    add_styled_table(["Giai đoạn / Tiến trình", "Nội dung công việc chi tiết"], project_steps, [2.0, 4.5])

    add_heading("3.2. Phân tích thiết kế hệ thống", level=2, bookmark="sec_3_2")
    add_heading("3.2.1. Tác nhân và ca sử dụng", level=3, bookmark="sec_3_2_1")
    add_p("Hệ thống xác định 3 tác nhân chính tham gia tương tác:")
    add_p("- Khách hàng (ROLE_USER): Người có nhu cầu tìm kiếm, thuê phòng trọ; theo dõi trạng thái yêu cầu thuê, đánh giá/bình luận, lưu tin bài đăng, quản lý danh sách phòng đã thuê và nhắn tin trực tiếp với chủ trọ.")
    add_p("- Người cho thuê / Chủ trọ (ROLE_RENTALER): Người trực tiếp quản lý và vận hành nhà trọ; đăng bài cho thuê, quản lý thông tin phòng, lập hợp đồng thuê phòng, ghi nhận và tính tiền điện nước hàng tháng, quản lý bảo trì tài sản, tiếp nhận yêu cầu từ khách và chat với khách thuê.")
    add_p("- Người quản trị (ROLE_ADMIN): Người có quyền hạn cao nhất kiểm soát toàn bộ hệ thống; kiểm duyệt tin đăng phòng mới, gỡ bỏ các bài đăng vi phạm, quản lý và phân quyền tài khoản, khóa tài khoản vi phạm và gửi email thông báo.")

    add_heading("3.2.2. Danh sách Use Case", level=3, bookmark="sec_3_2_2")
    add_p("Các ca sử dụng cốt lõi của hệ thống:")
    add_p("1. Quản lý phòng và bài đăng (Thêm mới, chỉnh sửa, ẩn bài, phê duyệt, gỡ bài đăng).")
    add_p("2. Quản lý hợp đồng (Lập hợp đồng, cập nhật hạn hợp đồng, tải tệp scan hợp đồng, xuất hóa đơn/thủ tục trả phòng).")
    add_p("3. Quản lý chỉ số điện nước (Ghi nhận số điện/nước hàng tháng, tính tổng tiền, cập nhật trạng thái thanh toán).")
    add_p("4. Quản lý bảo trì (Tạo phiếu bảo trì, ghi nhận chi phí sửa chữa, tải hóa đơn đính kèm).")
    add_p("5. Quản lý yêu cầu (Khách tạo yêu cầu lắp đặt/hỗ trợ, chủ trọ xử lý yêu cầu và xuất hóa đơn dịch vụ).")
    add_p("6. Quản lý tài khoản (Đăng nhập, đăng ký, xác thực email, khóa tài khoản, phân quyền).")
    add_p("7. Nhắn tin trực tuyến (Trò chuyện trực tiếp qua WebSocket/STOMP giữa khách và chủ trọ).")
    add_p("8. Lập thống kê báo cáo (Thống kê doanh thu theo tháng, cơ cấu chi phí, số lượng phòng trống/đã thuê).")
    add_p("9. Đánh giá và bình luận (Khách thuê chấm điểm sao và gửi nhận xét về phòng trọ).")

    add_heading("3.2.3. Đặc tả chi tiết các ca sử dụng", level=3, bookmark="sec_3_2_3")
    add_p("- Quản lý phòng và bài đăng: Chủ trọ đăng nhập vào hệ thống, chọn chức năng Thêm phòng. Nhập đầy đủ tiêu đề, mô tả, giá phòng, đơn giá nước, điện chung, internet, chọn tọa độ vị trí trên Google Maps và tải ảnh phòng trọ. Phòng mới tạo sẽ ở trạng thái 'Chờ duyệt'. Quản trị viên kiểm tra nội dung và click 'Phê duyệt' để bài đăng hiển thị công khai. Nếu bài đăng vi phạm, Admin có quyền 'Gỡ bài đăng'. Chủ trọ có thể 'Ẩn bài' khi phòng đã đầy.")
    add_p("- Quản lý hợp đồng: Khi có khách thỏa thuận thuê, chủ trọ chọn chức năng 'Lập hợp đồng', nhập họ tên người thuê, số CMND/CCCD, số điện thoại, số lượng người ở, ngày bắt đầu, ngày hết hạn và tải tệp hợp đồng đính kèm. Khi khách trả phòng, chủ trọ bấm 'Xuất hóa đơn trả phòng' để thanh lý hợp đồng và cập nhật trạng thái phòng về 'Còn trống'.")
    add_p("- Quản lý chỉ số Điện - Nước: Định kỳ hàng tháng, chủ trọ chọn phòng cần ghi số, nhập chỉ số điện cũ/mới, số khối nước cũ/mới và đơn giá. Hệ thống tự động tính thành tiền điện, tiền nước và tổng số tiền phải trả, cập nhật trạng thái 'Đã thanh toán' hoặc 'Chưa thanh toán'.")
    add_p("- Quản lý bảo trì: Khi phòng phát sinh hư hỏng trang thiết bị, chủ trọ tạo 'Phiếu bảo trì', nhập ngày sửa chữa, nội dung hư hỏng, chi phí phát sinh và tải hóa đơn sửa chữa lưu vào CSDL.")
    add_p("- Quản lý yêu cầu: Khách thuê gửi yêu cầu sửa chữa hoặc lắp thêm thiết bị qua form 'Gửi yêu cầu'. Chủ trọ nhận được thông báo, tiến hành xử lý yêu cầu và xuất hóa đơn thanh toán cho khách.")
    add_p("- Quản lý tài khoản: Quản trị viên theo dõi danh sách tất cả người dùng trong hệ thống; có thể khóa/mở khóa tài khoản nếu vi phạm điều khoản; thực hiện phân quyền (ROLE_USER, ROLE_RENTALER, ROLE_ADMIN).")
    add_p("- Lập thống kê: Chủ trọ xem biểu đồ doanh thu theo từng tháng trong năm, cơ cấu thu từ tiền phòng và tiền dịch vụ phụ trợ (điện, nước, net). Admin xem thống kê tổng số người dùng, số phòng trên toàn hệ thống.")
    add_p("- Đánh giá và bình luận: Khách thuê sau khi trải nghiệm có thể gửi đánh giá từ 1 đến 5 sao và bình luận nhận xét chất lượng phòng.")

    add_heading("3.3. Biểu đồ ca sử dụng (Use Case Diagram)", level=2, bookmark="sec_3_3")
    add_image_placeholder("Hình 3.1", "Biểu đồ ca sử dụng của người cho thuê")
    add_image_placeholder("Hình 3.2", "Biểu đồ ca sử dụng của quản trị viên")
    add_image_placeholder("Hình 3.3", "Biểu đồ phân quyền của quản trị viên")
    add_image_placeholder("Hình 3.4", "Biểu đồ ca sử dụng thuê phòng")
    add_image_placeholder("Hình 3.5", "Biểu đồ ca sử dụng đánh giá và bình luận")

    add_heading("3.4. Biểu đồ trình tự (Sequence Diagram)", level=2, bookmark="sec_3_4")
    add_heading("3.4.1. Quản lý thuê phòng", level=3, bookmark="sec_3_4_1")
    add_image_placeholder("Hình 3.6", "Biểu đồ trình tự thuê phòng")

    add_heading("3.4.2. Quản lý hợp đồng", level=3, bookmark="sec_3_4_2")
    add_image_placeholder("Hình 3.7", "Biểu đồ trình tự hợp đồng")

    add_heading("3.4.3. Quản lý phòng, bài đăng", level=3, bookmark="sec_3_4_3")
    add_image_placeholder("Hình 3.8", "Biểu đồ trình tự phòng, bài đăng")

    add_heading("3.4.4. Quản lý bảo trì", level=3, bookmark="sec_3_4_4")
    add_image_placeholder("Hình 3.9", "Biểu đồ trình tự về phiếu bảo trì")

    add_heading("3.4.5. Quản lý yêu cầu", level=3, bookmark="sec_3_4_5")
    add_image_placeholder("Hình 3.10", "Biểu đồ trình tự phiếu yêu cầu")

    add_heading("3.4.6. Quản lý tài khoản", level=3, bookmark="sec_3_4_6")
    add_image_placeholder("Hình 3.11", "Biểu đồ trình tự quản lý tài khoản")

    add_heading("3.4.7. Chức năng gỡ và phê duyệt bài đăng", level=3, bookmark="sec_3_4_7")
    add_image_placeholder("Hình 3.12", "Biểu đồ trình tự phê duyệt và gỡ phòng, bài đăng")

    add_heading("3.5. Biểu đồ trạng thái (State Diagram)", level=2, bookmark="sec_3_5")
    add_heading("3.5.1. Trạng thái Tài khoản", level=3, bookmark="sec_3_5_1")
    add_image_placeholder("Hình 3.13", "Biểu đồ trạng thái của tài khoản")

    add_heading("3.5.2. Trạng thái Phòng", level=3, bookmark="sec_3_5_2")
    add_image_placeholder("Hình 3.14", "Biểu đồ trạng thái của phòng")

    add_heading("3.5.3. Trạng thái Bài đăng của phòng", level=3, bookmark="sec_3_5_3")
    add_image_placeholder("Hình 3.15", "Biểu đồ trạng thái của bài đăng phòng")

    add_heading("3.5.4. Trạng thái Xử lý yêu cầu", level=3, bookmark="sec_3_5_4")
    add_image_placeholder("Hình 3.16", "Biểu đồ trạng thái của xử lý yêu cầu")

    add_heading("3.6. Biểu đồ hoạt động (Activity Diagram)", level=2, bookmark="sec_3_6")
    add_p("Lưu ý: Tại các điểm rẽ nhánh điều kiện (Decision Diamond), luồng xử lý được phân định rõ ràng giữa trường hợp dữ liệu [Hợp lệ] và [Không hợp lệ / Thiếu thông tin].")
    
    add_heading("3.6.1. Chức năng đăng nhập", level=3, bookmark="sec_3_6_1")
    add_image_placeholder("Hình 3.17", "Biểu đồ hoạt động của đăng nhập")

    add_heading("3.6.2. Chức năng đăng ký", level=3, bookmark="sec_3_6_2")
    add_image_placeholder("Hình 3.18", "Biểu đồ hoạt động của đăng ký")

    add_heading("3.6.3. Chức năng lập hợp đồng thuê phòng", level=3, bookmark="sec_3_6_3")
    add_image_placeholder("Hình 3.19", "Biểu đồ hoạt động của hợp đồng")

    add_heading("3.6.4. Chức năng lập phiếu bảo trì", level=3, bookmark="sec_3_6_4")
    add_image_placeholder("Hình 3.20", "Biểu đồ hoạt động của bảo trì")

    add_heading("3.6.5. Chức năng lập phiếu yêu cầu", level=3, bookmark="sec_3_6_5")
    add_image_placeholder("Hình 3.21", "Biểu đồ hoạt động của phiếu yêu cầu")

    add_heading("3.6.6. Chức năng thêm phòng, bài đăng", level=3, bookmark="sec_3_6_6")
    add_image_placeholder("Hình 3.22", "Biểu đồ hoạt động của chức năng thêm phòng, bài đăng")

    add_heading("3.6.7. Chức năng phê duyệt phòng", level=3, bookmark="sec_3_6_7")
    add_image_placeholder("Hình 3.23", "Biểu đồ hoạt động của phê duyệt bài đăng")

    add_heading("3.6.8. Chức năng gỡ phòng, bài đăng", level=3, bookmark="sec_3_6_8")
    add_image_placeholder("Hình 3.24", "Biểu đồ hoạt động của gỡ bài đăng")

    add_heading("3.6.9. Chức năng phân quyền", level=3, bookmark="sec_3_6_9")
    add_image_placeholder("Hình 3.25", "Biểu đồ hoạt động của chức năng phân quyền")

    add_heading("3.6.10. Chức năng bình luận", level=3, bookmark="sec_3_6_10")
    add_image_placeholder("Hình 3.26", "Biểu đồ hoạt động của chức năng bình luận")

    add_heading("3.6.11. Chức năng đánh giá", level=3, bookmark="sec_3_6_11")
    add_image_placeholder("Hình 3.27", "Biểu đồ hoạt động của chức năng đánh giá")

    add_heading("3.7. Biểu đồ lớp chi tiết (Class Diagram)", level=2, bookmark="sec_3_7")
    add_image_placeholder("Hình 3.28", "Biểu đồ lớp chi tiết")

    add_heading("3.8. Biểu đồ thành phần (Component Diagram)", level=2, bookmark="sec_3_8")
    add_heading("3.8.1. Người quản trị", level=3, bookmark="sec_3_8_1")
    add_image_placeholder("Hình 3.29", "Biểu đồ thành phần người quản trị")

    add_heading("3.8.2. Người cho thuê", level=3, bookmark="sec_3_8_2")
    add_image_placeholder("Hình 3.30", "Biểu đồ thành phần người cho thuê")

    add_heading("3.9. Thiết kế cơ sở dữ liệu", level=2, bookmark="sec_3_9")
    add_heading("3.9.1. Cơ sở dữ liệu và mối quan hệ (ERD)", level=3, bookmark="sec_3_9_1")
    add_image_placeholder("Hình 3.31", "Cơ sở dữ liệu và mối quan hệ (ERD)")

    add_heading("3.9.2. Mô tả chi tiết các bảng dữ liệu", level=3, bookmark="sec_3_9_2")
    add_p("Dưới đây là đặc tả chi tiết cấu trúc các bảng trong cơ sở dữ liệu quan hệ MySQL:")

    # Table 3.2: User
    add_p("Bảng 3.2: Bảng User (Người dùng hệ thống)", bold=True)
    user_cols = [
        ("id", "BIGINT", "Mã định danh duy nhất của người dùng", "Khóa chính, Tự tăng"),
        ("name", "VARCHAR(255)", "Họ và tên người dùng", "Bắt buộc"),
        ("email", "VARCHAR(255)", "Địa chỉ thư điện tử", "Bắt buộc, Duy nhất"),
        ("password", "VARCHAR(255)", "Mật khẩu đã được mã hóa BCrypt", "Cho tài khoản nội bộ"),
        ("phone", "VARCHAR(20)", "Số điện thoại liên hệ", "Duy nhất"),
        ("address", "VARCHAR(255)", "Địa chỉ cư trú", ""),
        ("image_url", "VARCHAR(255)", "Đường dẫn ảnh đại diện", ""),
        ("is_locked", "BIT(1)", "Trạng thái khóa tài khoản (0: Hoạt động, 1: Khóa)", ""),
        ("is_confirmed", "BIT(1)", "Trạng thái xác thực email đăng ký", ""),
        ("provider", "VARCHAR(50)", "Nhà cung cấp xác thực (LOCAL, GOOGLE, FACEBOOK)", ""),
        ("provider_id", "VARCHAR(255)", "Mã định danh phía Google / Facebook", ""),
        ("facebook_url", "VARCHAR(255)", "Đường dẫn trang Facebook cá nhân", ""),
        ("zalo_url", "VARCHAR(255)", "Đường dẫn Zalo liên hệ", ""),
        ("created_at", "DATETIME", "Thời điểm tạo tài khoản", "Tự động"),
        ("updated_at", "DATETIME", "Thời điểm cập nhật thông tin", "Tự động")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], user_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.3: Role
    add_p("Bảng 3.3: Bảng Role (Vai trò và quyền hạn)", bold=True)
    role_cols = [
        ("id", "BIGINT", "Mã định danh quyền", "Khóa chính, Tự tăng"),
        ("name", "VARCHAR(60)", "Tên quyền (ROLE_USER, ROLE_RENTALER, ROLE_ADMIN)", "Duy nhất")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], role_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.4: Room
    add_p("Bảng 3.4: Bảng Room (Phòng trọ và tin đăng cho thuê)", bold=True)
    room_cols = [
        ("id", "BIGINT", "Mã định danh phòng", "Khóa chính, Tự tăng"),
        ("title", "VARCHAR(255)", "Tiêu đề bài đăng phòng trọ", "Bắt buộc"),
        ("description", "TEXT", "Mô tả chi tiết phòng trọ và tiện ích", ""),
        ("price", "DECIMAL(19,2)", "Giá thuê phòng hàng tháng (VNĐ)", "Bắt buộc"),
        ("water_cost", "DECIMAL(19,2)", "Đơn giá tiền nước hoặc giá cố định (VNĐ)", "Mặc định 0"),
        ("public_electric_cost", "DECIMAL(19,2)", "Tiền điện sinh hoạt chung / hành lang (VNĐ)", "Mặc định 0"),
        ("internet_cost", "DECIMAL(19,2)", "Tiền dịch vụ mạng Internet (VNĐ)", "Mặc định 0"),
        ("address", "VARCHAR(255)", "Địa chỉ cụ thể của phòng trọ", "Bắt buộc"),
        ("latitude", "DOUBLE", "Tọa độ vĩ độ (Google Maps)", ""),
        ("longitude", "DOUBLE", "Tọa độ kinh độ (Google Maps)", ""),
        ("status", "VARCHAR(50)", "Trạng thái phòng (ROOM_EMPTY, ROOM_HIRED)", ""),
        ("is_locked", "VARCHAR(50)", "Trạng thái khóa tin (ENABLE, DISABLE)", ""),
        ("is_approve", "BIT(1)", "Admin đã duyệt bài đăng hay chưa (0: Chưa, 1: Duyệt)", ""),
        ("is_remove", "BIT(1)", "Admin đã gỡ bài đăng hay chưa (0: Bình thường, 1: Gỡ)", ""),
        ("user_id", "BIGINT", "Mã chủ phòng trọ (Rentaler)", "Khóa phụ"),
        ("category_id", "BIGINT", "Mã danh mục phòng trọ", "Khóa phụ"),
        ("location_id", "BIGINT", "Mã khu vực / Thành phố", "Khóa phụ"),
        ("created_at", "DATETIME", "Thời điểm tạo phòng", "Tự động"),
        ("updated_at", "DATETIME", "Thời điểm cập nhật", "Tự động"),
        ("created_by", "VARCHAR(255)", "Tên người tạo", ""),
        ("updated_by", "VARCHAR(255)", "Tên người cập nhật", "")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], room_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.5: Rate
    add_p("Bảng 3.5: Bảng Rate (Đánh giá xếp hạng sao)", bold=True)
    rate_cols = [
        ("id", "BIGINT", "Mã định danh đánh giá", "Khóa chính, Tự tăng"),
        ("rating", "DOUBLE", "Số điểm sao đánh giá (từ 1.0 đến 5.0 sao)", "Bắt buộc"),
        ("description", "VARCHAR(255)", "Mô tả tóm tắt đánh giá", ""),
        ("room_id", "BIGINT", "Mã phòng được đánh giá", "Khóa phụ"),
        ("user_id", "BIGINT", "Mã khách hàng thực hiện đánh giá", "Khóa phụ"),
        ("created_at", "DATETIME", "Thời gian đánh giá", "Tự động"),
        ("updated_at", "DATETIME", "Thời gian cập nhật", "Tự động")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], rate_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.6: Follow
    add_p("Bảng 3.6: Bảng Follow (Theo dõi chủ nhà trọ)", bold=True)
    follow_cols = [
        ("id", "BIGINT", "Mã định danh theo dõi", "Khóa chính, Tự tăng"),
        ("customer_id", "BIGINT", "Mã khách hàng theo dõi", "Khóa phụ"),
        ("rentaler_id", "BIGINT", "Mã chủ trọ được theo dõi", "Khóa phụ"),
        ("created_at", "DATETIME", "Thời điểm theo dõi", "Tự động"),
        ("updated_at", "DATETIME", "Thời điểm cập nhật", "Tự động")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], follow_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.7: BlogStore
    add_p("Bảng 3.7: Bảng BlogStore (Lưu tin bài đăng yêu thích)", bold=True)
    blog_cols = [
        ("id", "BIGINT", "Mã lưu tin", "Khóa chính, Tự tăng"),
        ("room_id", "BIGINT", "Mã phòng trọ được lưu", "Khóa phụ"),
        ("user_id", "BIGINT", "Mã người dùng lưu tin", "Khóa phụ"),
        ("created_at", "DATETIME", "Thời điểm lưu tin", "Tự động"),
        ("updated_at", "DATETIME", "Thời điểm cập nhật", "Tự động")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], blog_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.8: Message
    add_p("Bảng 3.8: Bảng Message (Cuộc hội thoại nhắn tin)", bold=True)
    msg_cols = [
        ("id", "BIGINT", "Mã định danh cuộc trò chuyện", "Khóa chính, Tự tăng"),
        ("user1", "BIGINT", "Mã người dùng thứ nhất (Người gửi khởi tạo)", "Khóa phụ"),
        ("user2", "BIGINT", "Mã người dùng thứ hai (Người nhận)", "Khóa phụ")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], msg_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.9: MessageChat
    add_p("Bảng 3.9: Bảng MessageChat (Nội dung tin nhắn chi tiết)", bold=True)
    chat_cols = [
        ("id", "BIGINT", "Mã tin nhắn", "Khóa chính, Tự tăng"),
        ("content", "TEXT", "Nội dung văn bản tin nhắn", "Bắt buộc"),
        ("is_read", "BIT(1)", "Trạng thái đối phương đã đọc hay chưa", ""),
        ("send_by", "BIT(1)", "Chiều gửi tin nhắn (0: user1 gửi, 1: user2 gửi)", ""),
        ("sent_at", "DATETIME", "Thời điểm gửi tin nhắn", "Tự động"),
        ("message_id", "BIGINT", "Mã cuộc trò chuyện tương ứng", "Khóa phụ")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], chat_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.10: Asset
    add_p("Bảng 3.10: Bảng Asset (Tài sản và trang thiết bị phòng trọ)", bold=True)
    asset_cols = [
        ("id", "BIGINT", "Mã tài sản", "Khóa chính, Tự tăng"),
        ("name", "VARCHAR(255)", "Tên tài sản / Thiết bị (Điều hòa, Tủ lạnh, Giường...)", "Bắt buộc"),
        ("number", "INT", "Số lượng trang thiết bị bàn giao trong phòng", ""),
        ("room_id", "BIGINT", "Mã phòng trọ sở hữu tài sản", "Khóa phụ")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], asset_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.11: Comment
    add_p("Bảng 3.11: Bảng Comment (Bình luận nhận xét)", bold=True)
    comment_cols = [
        ("id", "BIGINT", "Mã bình luận", "Khóa chính, Tự tăng"),
        ("content", "TEXT", "Nội dung nhận xét của khách thuê", "Bắt buộc"),
        ("room_id", "BIGINT", "Mã phòng được bình luận", "Khóa phụ"),
        ("user_id", "BIGINT", "Mã người dùng viết bình luận", "Khóa phụ"),
        ("rate_id", "BIGINT", "Mã đánh giá sao liên kết", "Khóa phụ"),
        ("created_at", "DATETIME", "Thời điểm bình luận", "Tự động"),
        ("updated_at", "DATETIME", "Thời điểm cập nhật", "Tự động"),
        ("created_by", "VARCHAR(255)", "Người viết bình luận", ""),
        ("updated_by", "VARCHAR(255)", "Người chỉnh sửa bình luận", "")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], comment_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.12: Category
    add_p("Bảng 3.12: Bảng Category (Danh mục phân loại phòng)", bold=True)
    cat_cols = [
        ("id", "BIGINT", "Mã danh mục", "Khóa chính, Tự tăng"),
        ("name", "VARCHAR(255)", "Tên danh mục (Phòng trọ, Chung cư mini, Nhà nguyên căn)", "Bắt buộc")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], cat_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.13: Location
    add_p("Bảng 3.13: Bảng Location (Địa điểm / Tỉnh thành)", bold=True)
    loc_cols = [
        ("id", "BIGINT", "Mã khu vực", "Khóa chính, Tự tăng"),
        ("city_name", "VARCHAR(255)", "Tên tỉnh thành phố (Hà Nội, TP.HCM, Đà Nẵng...)", "Bắt buộc")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], loc_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.14: Maintenance
    add_p("Bảng 3.14: Bảng Maintenance (Phiếu bảo trì sửa chữa)", bold=True)
    maint_cols = [
        ("id", "BIGINT", "Mã phiếu bảo trì", "Khóa chính, Tự tăng"),
        ("maintenance_date", "DATETIME", "Thời điểm tiến hành bảo trì sửa chữa", "Bắt buộc"),
        ("price", "DECIMAL(19,2)", "Chi phí bảo trì sửa chữa (VNĐ)", ""),
        ("files", "VARCHAR(255)", "Đường dẫn tệp hóa đơn hoặc biên bản sửa chữa", ""),
        ("room_id", "BIGINT", "Mã phòng được bảo trì", "Khóa phụ"),
        ("created_at", "DATETIME", "Thời điểm tạo phiếu", "Tự động"),
        ("updated_at", "DATETIME", "Thời điểm cập nhật", "Tự động"),
        ("created_by", "VARCHAR(255)", "Người lập phiếu", ""),
        ("updated_by", "VARCHAR(255)", "Người cập nhật", "")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], maint_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.15: Request
    add_p("Bảng 3.15: Bảng Request (Phiếu yêu cầu từ khách thuê)", bold=True)
    req_cols = [
        ("id", "BIGINT", "Mã phiếu yêu cầu", "Khóa chính, Tự tăng"),
        ("name", "VARCHAR(255)", "Họ tên người thuê gửi yêu cầu", "Bắt buộc"),
        ("phone_number", "VARCHAR(20)", "Số điện thoại liên hệ của người thuê", "Bắt buộc"),
        ("description", "TEXT", "Nội dung chi tiết yêu cầu (sửa điện nước, lắp thêm đồ)", "Bắt buộc"),
        ("is_answer", "BIT(1)", "Trạng thái chủ trọ đã xử lý yêu cầu hay chưa", "Mặc định 0"),
        ("room_id", "BIGINT", "Mã phòng trọ tương ứng", "Khóa phụ")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], req_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.16: Contract
    add_p("Bảng 3.16: Bảng Contract (Hợp đồng thuê phòng trọ)", bold=True)
    contract_cols = [
        ("id", "BIGINT", "Mã hợp đồng thuê", "Khóa chính, Tự tăng"),
        ("name", "VARCHAR(255)", "Tên hợp đồng", "Bắt buộc"),
        ("name_of_rent", "VARCHAR(255)", "Họ tên người đại diện đứng tên thuê phòng", "Bắt buộc"),
        ("phone", "VARCHAR(20)", "Số điện thoại người đại diện thuê", "Bắt buộc"),
        ("num_of_people", "INT", "Số lượng người ở thực tế trong phòng", "Bắt buộc"),
        ("deadline_contract", "DATETIME", "Thời hạn kết thúc hợp đồng thuê", "Bắt buộc"),
        ("files", "VARCHAR(255)", "Đường dẫn file scan hợp đồng đính kèm", ""),
        ("room_id", "BIGINT", "Mã phòng được thuê", "Khóa phụ"),
        ("created_at", "DATETIME", "Thời điểm ký hợp đồng", "Tự động"),
        ("updated_at", "DATETIME", "Thời điểm cập nhật", "Tự động"),
        ("created_by", "VARCHAR(255)", "Người lập hợp đồng", ""),
        ("updated_by", "VARCHAR(255)", "Người cập nhật", "")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], contract_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.17: RoomMedia
    add_p("Bảng 3.17: Bảng RoomMedia (Tệp hình ảnh phòng trọ)", bold=True)
    media_cols = [
        ("id", "BIGINT", "Mã ảnh phòng", "Khóa chính, Tự tăng"),
        ("files", "VARCHAR(255)", "Đường dẫn tệp hình ảnh lưu trữ trên server", "Bắt buộc"),
        ("room_id", "BIGINT", "Mã phòng sở hữu hình ảnh", "Khóa phụ")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], media_cols, [1.5, 1.5, 2.5, 1.5])

    # Table 3.18: ElectricAndWater
    add_p("Bảng 3.18: Bảng ElectricAndWater (Quản lý chỉ số Điện - Nước theo tháng)", bold=True)
    ew_cols = [
        ("id", "BIGINT", "Mã bản ghi điện nước", "Khóa chính, Tự tăng"),
        ("name", "VARCHAR(255)", "Tiêu đề ghi nhận (ví dụ: Hóa đơn điện nước T10/2026)", "Bắt buộc"),
        ("month", "INT", "Tháng tính toán (1 - 12)", "Bắt buộc"),
        ("last_month_number_of_electric", "INT", "Chỉ số công tơ điện tháng trước (kWh)", "Bắt buộc"),
        ("this_month_number_of_electric", "INT", "Chỉ số công tơ điện tháng này (kWh)", "Bắt buộc"),
        ("money_each_number_of_electric", "DECIMAL(19,2)", "Đơn giá tiền điện trên 1 số (VNĐ/kWh)", "Bắt buộc"),
        ("total_money_of_electric", "DECIMAL(19,2)", "Tổng tiền điện trong tháng (VNĐ)", "Tự động tính"),
        ("last_month_block_of_water", "INT", "Chỉ số đồng hồ nước tháng trước (m³)", "Bắt buộc"),
        ("this_month_block_of_water", "INT", "Chỉ số đồng hồ nước tháng này (m³)", "Bắt buộc"),
        ("money_each_block_of_water", "DECIMAL(19,2)", "Đơn giá tiền nước trên 1 khối (VNĐ/m³)", "Bắt buộc"),
        ("total_money_of_water", "DECIMAL(19,2)", "Tổng tiền nước trong tháng (VNĐ)", "Tự động tính"),
        ("paid", "BIT(1)", "Trạng thái thanh toán (0: Chưa thanh toán, 1: Đã thanh toán)", "Mặc định 0"),
        ("room_id", "BIGINT", "Mã phòng trọ thanh toán", "Khóa phụ")
    ]
    add_styled_table(["Tên trường", "Kiểu dữ liệu", "Mô tả trường", "Ghi chú"], ew_cols, [1.5, 1.5, 2.5, 1.5])

    doc.add_page_break()

    # =========================================================================
    # CHƯƠNG 4: KẾT QUẢ CÀI ĐẶT CHƯƠNG TRÌNH
    # =========================================================================
    add_heading("CHƯƠNG 4: KẾT QUẢ CÀI ĐẶT CHƯƠNG TRÌNH", level=1, bookmark="sec_ch4")
    add_p("Hệ thống website quản lý và cho thuê nhà trọ đã được cài đặt và vận hành hoàn chỉnh với đầy đủ các phân hệ chức năng tương ứng cho 3 nhóm tác nhân. Dưới đây là hình ảnh giao diện thực tế cùng mô tả chi tiết luồng hoạt động của từng chức năng:")

    add_heading("4.1. Chức năng dành cho người thuê (Khách hàng)", level=2, bookmark="sec_4_1")
    add_p("1. Trang chủ và Danh sách phòng trọ:")
    add_p("Giao diện trang chủ hiển thị danh sách các phòng trọ đang cho thuê kèm bộ lọc nâng cao theo: Tên phòng/từ khóa, Mức giá mong muốn, Danh mục phòng trọ (Phòng trọ sinh viên, Chung cư mini, Nhà nguyên căn). Hỗ trợ phân trang tiện lợi.")
    add_image_placeholder("Hình 4.1", "Giao diện Trang chủ và Danh sách phòng trọ")

    add_p("2. Trang Chi tiết phòng trọ và Tích hợp Bản đồ Google Maps:")
    add_p("Hiển thị thư viện ảnh phòng dạng trình chiếu Swiper mượt mà, đầy đủ thông tin: giá thuê hàng tháng, chi phí tiền nước, tiền điện chung, internet, danh sách tài sản tiện nghi trong phòng. Bản đồ Google Maps hiển thị vị trí chính xác của phòng trọ giúp khách dễ dàng tìm đường.")
    add_image_placeholder("Hình 4.2", "Giao diện Chi tiết phòng trọ và Bản đồ Google Maps")

    add_p("3. Gửi yêu cầu đặt thuê phòng và Liên hệ:")
    add_p("Khách hàng có thể gửi yêu cầu đặt hẹn xem phòng hoặc đăng ký thuê trực tiếp với chủ trọ; đồng thời hệ thống hỗ trợ form gửi email liên hệ báo tin tức thì tới chủ nhà.")
    add_image_placeholder("Hình 4.3", "Giao diện Gửi yêu cầu đặt thuê phòng")

    add_p("4. Đánh giá xếp hạng sao và Bình luận:")
    add_p("Khách thuê có thể chấm điểm sao từ 1 đến 5 sao và viết nhận xét chi tiết về chất lượng phòng, thái độ phục vụ của chủ nhà.")
    add_image_placeholder("Hình 4.4", "Giao diện Đánh giá, bình luận phòng trọ")

    add_p("5. Nhắn tin trực tuyến (Chat Realtime):")
    add_p("Tích hợp giao diện trò chuyện trực tiếp với chủ trọ qua giao thức WebSocket và STOMP, thông báo tin nhắn mới tức thời mà không cần làm mới trang.")
    add_image_placeholder("Hình 4.5", "Giao diện Nhắn tin trực tuyến (Chat Realtime)")

    add_p("6. Tra cứu phòng đã thuê:")
    add_p("Khách hàng tra cứu nhanh chóng danh sách các phòng mình đã hoặc đang thuê theo số điện thoại cá nhân.")
    add_image_placeholder("Hình 4.6", "Giao diện Quản lý danh sách phòng đã thuê")

    add_heading("4.2. Chức năng dành cho người cho thuê (Chủ trọ)", level=2, bookmark="sec_4_2")
    add_p("1. Dashboard thống kê doanh thu và chi phí:")
    add_p("Hiển thị tổng quan các chỉ số trọng yếu: Tổng số phòng, số phòng đang có khách, số phòng còn trống, số lượng khách thuê. Hệ thống vẽ biểu đồ cột (BarChart) biểu diễn doanh thu từng tháng trong năm và biểu đồ tròn phân tích cơ cấu chi phí dịch vụ.")
    add_image_placeholder("Hình 4.7", "Giao diện Dashboard thống kê doanh thu và chi phí theo tháng")

    add_p("2. Quản lý danh sách phòng và Đăng phòng mới:")
    add_p("Cho phép chủ trọ đăng tin phòng mới, đính kèm nhiều ảnh, chọn vị trí ghim tọa độ GPS trên bản đồ Google Maps và cập nhật bảng giá dịch vụ (điện chung, nước, net).")
    add_image_placeholder("Hình 4.8", "Giao diện Quản lý danh sách phòng và Thêm mới phòng")

    add_p("3. Quản lý hợp đồng thuê phòng:")
    add_p("Lập hợp đồng thuê mới cho khách; quản lý thời hạn kết thúc hợp đồng; tải tệp hợp đồng và xuất phiếu thanh lý trả phòng khi khách chuyển đi.")
    add_image_placeholder("Hình 4.9", "Giao diện Quản lý và Lập hợp đồng thuê phòng")

    add_p("4. Quản lý chỉ số Điện - Nước và tính tiền hàng tháng:")
    add_p("Ghi nhận số điện cũ/mới, số khối nước cũ/mới hàng tháng; tự động tính ra tổng tiền phải thanh toán; đánh dấu trạng thái Đã thu tiền / Chưa thu tiền.")
    add_image_placeholder("Hình 4.10", "Giao diện Quản lý chỉ số Điện - Nước và tính tiền hàng tháng")

    add_p("5. Quản lý yêu cầu từ khách thuê và Xuất hóa đơn:")
    add_p("Tiếp nhận danh sách yêu cầu hỗ trợ hoặc lắp đặt đồ đạc từ khách; cập nhật trạng thái đã xử lý và xuất hóa đơn thu tiền gửi khách.")
    add_image_placeholder("Hình 4.11", "Giao diện Quản lý yêu cầu từ khách thuê và xuất hóa đơn")

    add_p("6. Quản lý bảo trì phòng:")
    add_p("Theo dõi lịch sử sửa chữa trang thiết bị hư hỏng, cập nhật kinh phí bảo trì và hóa đơn thanh toán.")
    add_image_placeholder("Hình 4.12", "Giao diện Quản lý bảo trì và chi phí bảo trì phòng")

    add_heading("4.3. Chức năng dành cho quản trị viên (Admin)", level=2, bookmark="sec_4_3")
    add_p("1. Dashboard quản trị toàn hệ thống:")
    add_p("Cung cấp các con số thống kê tổng hợp: Tổng số tài khoản, số phòng đã duyệt, số tin đăng đang chờ kiểm duyệt trên toàn hệ thống.")
    add_image_placeholder("Hình 4.13", "Giao diện Dashboard quản trị hệ thống Admin")

    add_p("2. Kiểm duyệt và Gỡ bài đăng phòng trọ:")
    add_p("Admin xem xét chi tiết nội dung tin đăng; duyệt bài để hiển thị công khai trên website hoặc gỡ bỏ các bài đăng sai sự thật, vi phạm điều khoản.")
    add_image_placeholder("Hình 4.14", "Giao diện Phê duyệt và Gỡ bỏ bài đăng phòng trọ")

    add_p("3. Quản lý tài khoản và Phân quyền người dùng:")
    add_p("Xem danh sách toàn bộ người dùng; thực hiện phân quyền (ROLE_USER, ROLE_RENTALER, ROLE_ADMIN); khóa các tài khoản vi phạm quy định.")
    add_image_placeholder("Hình 4.15", "Giao diện Quản lý tài khoản và Phân quyền người dùng")

    add_p("4. Gửi email thông báo từ hệ thống:")
    add_p("Cho phép Admin gửi email cảnh báo hoặc thông báo trực tiếp đến từng chủ trọ hoặc khách hàng thông qua dịch vụ JavaMailSender.")
    add_image_placeholder("Hình 4.16", "Giao diện Gửi email thông báo tới người dùng")

    doc.add_page_break()

    # =========================================================================
    # KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN
    # =========================================================================
    add_heading("KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", level=1, bookmark="sec_conclusion")
    add_p("Đề tài 'Nghiên cứu xây dựng hệ thống website quản lý và cho thuê nhà trọ' đã hoàn thành các mục tiêu nghiên cứu và xây dựng thành công một giải pháp phần mềm hoàn chỉnh, giải quyết hiệu quả bài toán quản lý phòng trọ trong thực tiễn.")

    add_p("Kết quả đạt được của đề tài:", bold=True)
    add_p("1. Về mặt lý thuyết: Đã nghiên cứu và làm chủ kiến trúc phát triển ứng dụng Full-stack hiện đại: Spring Boot, Spring Security, JWT, WebSocket/STOMP, cơ sở dữ liệu MySQL và thư viện ReactJS.")
    add_p("2. Về mặt phân tích thiết kế: Xây dựng hoàn chỉnh hệ thống các biểu đồ chuẩn UML (Use Case, Sequence, State, Activity, Class, Component) và thiết kế cơ sở dữ liệu quan hệ tối ưu hóa.")
    add_p("3. Về mặt sản phẩm: Xây dựng thành công hệ thống website với đầy đủ các tính năng cho 3 vai trò:")
    add_p("   - Khách thuê: Tìm kiếm, lọc phòng, xem vị trí Google Maps, gửi yêu cầu đặt phòng, chat thời gian thực, đánh giá, lưu tin.")
    add_p("   - Chủ trọ: Quản lý tin đăng, hợp đồng, tự động hóa tính tiền điện nước theo tháng, quản lý bảo trì, xuất hóa đơn và theo dõi biểu đồ doanh thu.")
    add_p("   - Quản trị viên: Quản trị tài khoản, phân quyền, kiểm duyệt tin đăng và gửi email thông báo hệ thống.")
    add_p("4. Ứng dụng chạy ổn định, tương thích tốt trên các trình duyệt phổ biến như Google Chrome, Microsoft Edge, Cốc Cốc.")

    add_p("Hạn chế của đề tài:", bold=True)
    add_p("- Hệ thống hiện tại mới hỗ trợ lưu trữ tệp trên máy chủ cục bộ, chưa tích hợp dịch vụ lưu trữ đám mây (Cloud Storage như AWS S3 hoặc Cloudinary).")
    add_p("- Chưa tích hợp cổng thanh toán trực tuyến tự động (như VNPay, MoMo, ZaloPay) để khách thuê có thể thanh toán trực tiếp tiền cọc và tiền điện nước qua mã QR.")

    add_p("Hướng nghiên cứu và phát triển tiếp theo:", bold=True)
    add_p("1. Tích hợp cổng thanh toán điện tử VNPay / MoMo giúp tự động hóa khâu thu tiền phòng và gạch nợ hóa đơn điện nước tức thì.")
    add_p("2. Phát triển ứng dụng di động (Mobile App) trên nền tảng React Native hoặc Flutter để khách thuê và chủ trọ tiện theo dõi thông báo trên điện thoại thông minh.")
    add_p("3. Ứng dụng trí tuệ nhân tạo (AI/Machine Learning) trong việc gợi ý phòng trọ phù hợp dựa trên thói quen tìm kiếm và mức ngân sách của người dùng.")
    add_p("4. Nâng cấp hệ thống triển khai trên nền tảng Docker Container và đám mây (AWS/Google Cloud) để tăng khả năng chịu tải khi số lượng người dùng mở rộng.")

    doc.add_page_break()

    # =========================================================================
    # TÀI LIỆU THAM KHẢO
    # =========================================================================
    add_heading("TÀI LIỆU THAM KHẢO", level=1, bookmark="sec_ref")
    refs = [
        "1. Craig Walls (2018), Spring in Action, Fifth Edition, Manning Publications.",
        "2. Alex Banks, Eve Porcello (2020), Learning React: Modern Patterns for Developing React Apps, 2nd Edition, O'Reilly Media.",
        "3. Martin Fowler (2003), UML Distilled: A Brief Guide to the Standard Object Modeling Language, 3rd Edition, Addison-Wesley Professional.",
        "4. Paul DuBois (2013), MySQL: Developers Library, 5th Edition, Addison-Wesley Professional.",
        "5. Joshua Bloch (2018), Effective Java, 3rd Edition, Addison-Wesley Professional.",
        "6. Tài liệu chính thức của Spring Boot: https://spring.io/projects/spring-boot",
        "7. Tài liệu chính thức của ReactJS: https://react.dev",
        "8. Chuẩn tài liệu RFC 7519 về JSON Web Token (JWT): https://datatracker.ietf.org/doc/html/rfc7519",
        "9. Đặc tả giao thức STOMP qua WebSocket: https://stomp.github.io"
    ]
    for ref_text in refs:
        add_p(ref_text, space_after=8)

    # Save document
    doc.save(output_docx)
    print(f"BÁO CÁO ĐÃ ĐƯỢC TẠO THÀNH CÔNG TẠI: {output_docx}")
    return output_docx

if __name__ == "__main__":
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    path = build_complete_report()
    print("SUCCESS: File generated at", path)
