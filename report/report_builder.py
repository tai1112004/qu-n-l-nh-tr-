# -*- coding: utf-8 -*-
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import os

def create_report_document(output_path):
    doc = Document()

    # 1. Page Margins Setup (Standard Vietnamese Thesis: Top 2cm, Bottom 2cm, Left 3cm, Right 2cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.79)
        section.different_first_page_header_footer = True
        
        # Add Page Number in Footer (Centered)
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run = f_p.add_run()
        f_fld1 = OxmlElement('w:fldSimple')
        f_fld1.set(qn('w:instr'), 'PAGE')
        f_run._r.append(f_fld1)

    # 2. Settings: Enable updateFields on open
    update_fields = OxmlElement('w:updateFields')
    update_fields.set(qn('w:val'), 'true')
    doc.settings.element.append(update_fields)

    # 3. Base Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(13)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)
    style_normal.paragraph_format.line_spacing = 1.3
    style_normal.paragraph_format.space_after = Pt(6)

    # Heading Styles
    for level, size, bold, color in [
        (1, Pt(16), True, RGBColor(0, 51, 102)),
        (2, Pt(14), True, RGBColor(0, 51, 102)),
        (3, Pt(13), True, RGBColor(30, 41, 59))
    ]:
        h_style = doc.styles[f'Heading {level}']
        h_style.font.name = 'Times New Roman'
        h_style.font.size = size
        h_style.font.bold = bold
        h_style.font.color.rgb = color
        h_style.paragraph_format.line_spacing = 1.3
        h_style.paragraph_format.space_before = Pt(12)
        h_style.paragraph_format.space_after = Pt(6)

    # Helper Functions
    def add_custom_heading(text, level, bookmark):
        h = doc.add_heading(level=level)
        pPr = h._p.get_or_add_pPr()
        
        # Outline level for Word TOC recognition
        outlineLvl = parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="{level-1}"/>')
        pPr.append(outlineLvl)

        # Bookmark
        bm_id = (abs(hash(bookmark)) % 90000) + 1000
        bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bm_id}" w:name="{bookmark}"/>')
        h._p.append(bm_start)

        run = h.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0, 51, 102)
        elif level == 2:
            run.font.size = Pt(13.5)
            run.font.color.rgb = RGBColor(0, 51, 102)
        else:
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(30, 41, 59)

        bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bm_id}"/>')
        h._p.append(bm_end)
        return h

    def add_p(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.bold = bold
        run.italic = italic
        return p

    def add_image_placeholder(fig_label, caption, height_in=2.4):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(5.8)
        
        # Style cell: background color and dashed/solid border
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>
                <w:left w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>
                <w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>
                <w:right w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>
            </w:tcBorders>
        ''')
        cell._tc.get_or_add_tcPr().append(borders)
        
        # Content inside placeholder
        p_in = cell.paragraphs[0]
        p_in.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_in.paragraph_format.space_before = Pt(36)
        p_in.paragraph_format.space_after = Pt(36)
        
        r_icon = p_in.add_run("[ KHUNG CHÈN HÌNH ẢNH ]\n")
        r_icon.font.name = 'Times New Roman'
        r_icon.font.size = Pt(12)
        r_icon.bold = True
        r_icon.font.color.rgb = RGBColor(71, 85, 105)
        
        r_txt = p_in.add_run(f"(Vui lòng dán ảnh chụp hoặc sơ đồ cho: {caption})")
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(11)
        r_txt.italic = True
        r_txt.font.color.rgb = RGBColor(100, 116, 139)
        
        # Caption below table
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(f"{fig_label}: {caption}")
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(11.5)
        r_cap.italic = True
        r_cap.bold = True
        r_cap.font.color.rgb = RGBColor(15, 23, 42)

    def add_styled_table(headers, rows_data, col_widths=None):
        table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Style table borders
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr.append(tblBorders)

        # Header Row
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            cell = hdr_cells[i]
            cell.text = header_text
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(11)
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E3A8A"/>')
            cell._tc.get_or_add_tcPr().append(shd)

        # Data Rows
        for r_idx, row in enumerate(rows_data):
            row_cells = table.rows[r_idx + 1].cells
            bg_fill = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(row):
                cell = row_cells[c_idx]
                cell.text = str(cell_value)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                if c_idx == 0 or len(str(cell_value)) <= 8:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(11)
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_fill}"/>')
                cell._tc.get_or_add_tcPr().append(shd)

        # Apply Col Widths if specified
        if col_widths:
            for row in table.rows:
                for c_idx, w in enumerate(col_widths):
                    row.cells[c_idx].width = Inches(w)
                    
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(2)
        p_space.paragraph_format.space_after = Pt(6)

    print("Helper functions initialized.")
    return doc, add_custom_heading, add_p, add_image_placeholder, add_styled_table

print("create_report_document generator module ready.")
